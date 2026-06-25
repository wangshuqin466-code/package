import os
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = BASE_DIR / "google_trends_db.env"
FILE_PATTERN = "google_trends_wedding_formal_*.xlsx"


def load_env(path):
    config = {}
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            config[key.strip()] = value.strip().strip('"').strip("'")
    for key in [
        "DB_TYPE",
        "SQLITE_PATH",
        "DB_HOST",
        "DB_PORT",
        "DB_NAME",
        "DB_USER",
        "DB_PASSWORD",
    ]:
        if os.environ.get(key):
            config[key] = os.environ[key]
    return config


def latest_workbook():
    files = sorted(BASE_DIR.glob(FILE_PATTERN), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        raise FileNotFoundError(f"No workbook found matching {FILE_PATTERN} in {BASE_DIR}")
    return files[0]


def batch_id_from_path(path):
    match = re.search(r"(google_trends_wedding_formal_\d{8}_\d{6})", path.stem)
    if not match:
        raise ValueError(f"Cannot parse batch id from {path.name}")
    return match.group(1)


def rows_by_header(workbook_path, sheet_name):
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(cell or "").strip() for cell in next(rows)]
    except StopIteration:
        return []
    result = []
    for row in rows:
        item = {headers[idx]: row[idx] if idx < len(row) else None for idx in range(len(headers))}
        if any(value not in (None, "") for value in item.values()):
            result.append(item)
    return result


def iso_now():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def to_int(value):
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def sqlite_connect(config):
    db_path = Path(config.get("SQLITE_PATH", r"C:\tmp\google_trends_wedding_formal.db"))
    if not db_path.is_absolute():
        db_path = BASE_DIR / db_path
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def connect(config):
    db_type = config.get("DB_TYPE", "sqlite").lower()
    if db_type == "sqlite":
        return sqlite_connect(config), "sqlite"
    if db_type == "mysql":
        try:
            import pymysql
        except ImportError as exc:
            raise RuntimeError("MySQL needs pymysql. Install it with: python -m pip install pymysql") from exc
        return (
            pymysql.connect(
                host=config["DB_HOST"],
                port=int(config.get("DB_PORT", "3306")),
                user=config["DB_USER"],
                password=config["DB_PASSWORD"],
                database=config["DB_NAME"],
                charset="utf8mb4",
                autocommit=False,
            ),
            "mysql",
        )
    if db_type == "postgres":
        try:
            import psycopg
        except ImportError as exc:
            raise RuntimeError("PostgreSQL needs psycopg. Install it with: python -m pip install psycopg[binary]") from exc
        return (
            psycopg.connect(
                host=config["DB_HOST"],
                port=int(config.get("DB_PORT", "5432")),
                user=config["DB_USER"],
                password=config["DB_PASSWORD"],
                dbname=config["DB_NAME"],
            ),
            "postgres",
        )
    raise ValueError(f"Unsupported DB_TYPE: {db_type}")


def qmark(db_type):
    return "%s" if db_type in {"mysql", "postgres"} else "?"


def execute_many(conn, db_type, sql, rows):
    if not rows:
        return
    placeholder = qmark(db_type)
    if placeholder != "?":
        sql = sql.replace("?", placeholder)
    cur = conn.cursor()
    cur.executemany(sql, rows)


def execute(conn, db_type, sql, values=None):
    placeholder = qmark(db_type)
    if placeholder != "?":
        sql = sql.replace("?", placeholder)
    cur = conn.cursor()
    cur.execute(sql, values or [])
    return cur


def table_exists(conn, db_type, table_name):
    if db_type == "sqlite":
        cur = execute(
            conn,
            db_type,
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = ?",
            [table_name],
        )
    elif db_type == "mysql":
        cur = execute(conn, db_type, "SHOW TABLES LIKE ?", [table_name])
    else:
        cur = execute(
            conn,
            db_type,
            "SELECT table_name FROM information_schema.tables WHERE table_name = ?",
            [table_name],
        )
    return cur.fetchone() is not None


def create_tables(conn, db_type, include_related_topics):
    text_type = "VARCHAR(120)" if db_type == "mysql" else "TEXT"
    long_text_type = "TEXT"
    int_type = "INTEGER"
    ddl = [
        f"""
        CREATE TABLE IF NOT EXISTS crawl_runs (
            batch_id {text_type} PRIMARY KEY,
            source {text_type},
            workbook_path {long_text_type},
            status {text_type},
            started_at {text_type},
            finished_at {text_type},
            update_time {text_type}
        )
        """,
        f"""
        CREATE TABLE IF NOT EXISTS trend_timeseries (
            batch_id {text_type},
            geo_name {text_type},
            geo_code {text_type},
            timeframe_name {text_type},
            timeframe_param {text_type},
            date_label {text_type},
            timestamp {int_type},
            category {text_type},
            keyword {text_type},
            trend_value {int_type},
            formatted_value {text_type},
            source_url {long_text_type},
            update_time {text_type},
            PRIMARY KEY (batch_id, geo_code, timeframe_param, keyword, timestamp)
        )
        """,
        f"""
        CREATE TABLE IF NOT EXISTS trend_geo_heat (
            batch_id {text_type},
            geo_name {text_type},
            geo_code {text_type},
            timeframe_name {text_type},
            timeframe_param {text_type},
            subregion_name {text_type},
            subregion_code {text_type},
            category {text_type},
            keyword {text_type},
            trend_value {int_type},
            formatted_value {text_type},
            source_url {long_text_type},
            update_time {text_type},
            PRIMARY KEY (batch_id, geo_code, timeframe_param, keyword, subregion_code)
        )
        """,
        f"""
        CREATE TABLE IF NOT EXISTS trend_related_queries (
            batch_id {text_type},
            geo_name {text_type},
            geo_code {text_type},
            timeframe_name {text_type},
            timeframe_param {text_type},
            category {text_type},
            keyword {text_type},
            list_type {text_type},
            rank_no {int_type},
            query {text_type},
            trend_value {int_type},
            formatted_value {text_type},
            link {long_text_type},
            source_url {long_text_type},
            update_time {text_type},
            PRIMARY KEY (batch_id, geo_code, timeframe_param, keyword, list_type, rank_no, query)
        )
        """,
    ]
    if include_related_topics:
        ddl.append(
            f"""
            CREATE TABLE IF NOT EXISTS trend_related_topics (
                batch_id {text_type},
                geo_name {text_type},
                geo_code {text_type},
                timeframe_name {text_type},
                timeframe_param {text_type},
                category {text_type},
                keyword {text_type},
                list_type {text_type},
                rank_no {int_type},
                topic {text_type},
                topic_type {text_type},
                trend_value {int_type},
                formatted_value {text_type},
                link {long_text_type},
                source_url {long_text_type},
                update_time {text_type},
                PRIMARY KEY (batch_id, geo_code, timeframe_param, keyword, list_type, rank_no, topic)
            )
            """
        )
    for statement in ddl:
        execute(conn, db_type, statement)


def delete_batch(conn, db_type, batch_id):
    for table in [
        "trend_timeseries",
        "trend_geo_heat",
        "trend_related_queries",
        "trend_related_topics",
        "crawl_runs",
    ]:
        if table_exists(conn, db_type, table):
            execute(conn, db_type, f"DELETE FROM {table} WHERE batch_id = ?", [batch_id])


def import_workbook(conn, db_type, workbook_path):
    batch_id = batch_id_from_path(workbook_path)
    update_time = iso_now()
    timeseries = rows_by_header(workbook_path, "热度随时间变化趋势")
    geo_heat = rows_by_header(workbook_path, "搜索热度")
    related_queries = rows_by_header(workbook_path, "相关查询")
    related_topics = rows_by_header(workbook_path, "相关主题")
    timeseries = [row for row in timeseries if to_int(row.get("时间戳")) is not None]
    geo_heat = [
        row
        for row in geo_heat
        if row.get("子区域代码") not in (None, "") and to_int(row.get("热度值")) is not None
    ]
    related_queries = [
        row
        for row in related_queries
        if to_int(row.get("排名")) is not None and row.get("查询词") not in (None, "")
    ]
    related_topics = [
        row
        for row in related_topics
        if to_int(row.get("排名")) is not None and row.get("主题") not in (None, "")
    ]

    create_tables(conn, db_type, include_related_topics=bool(related_topics))
    delete_batch(conn, db_type, batch_id)

    execute(
        conn,
        db_type,
        """
        INSERT INTO crawl_runs
        (batch_id, source, workbook_path, status, started_at, finished_at, update_time)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        [batch_id, "google_trends", str(workbook_path), "importing", update_time, None, update_time],
    )

    execute_many(
        conn,
        db_type,
        """
        INSERT INTO trend_timeseries
        (batch_id, geo_name, geo_code, timeframe_name, timeframe_param, date_label, timestamp,
         category, keyword, trend_value, formatted_value, source_url, update_time)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                batch_id,
                row.get("区域"),
                row.get("区域代码"),
                row.get("时间维度"),
                row.get("时间参数"),
                row.get("日期"),
                to_int(row.get("时间戳")),
                row.get("分类"),
                row.get("关键词"),
                to_int(row.get("热度值")),
                row.get("格式化热度"),
                row.get("来源URL"),
                update_time,
            )
            for row in timeseries
        ],
    )

    execute_many(
        conn,
        db_type,
        """
        INSERT INTO trend_geo_heat
        (batch_id, geo_name, geo_code, timeframe_name, timeframe_param, subregion_name,
         subregion_code, category, keyword, trend_value, formatted_value, source_url, update_time)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                batch_id,
                row.get("区域"),
                row.get("区域代码"),
                row.get("时间维度"),
                row.get("时间参数"),
                row.get("子区域"),
                row.get("子区域代码"),
                row.get("分类"),
                row.get("关键词"),
                to_int(row.get("热度值")),
                row.get("格式化热度"),
                row.get("来源URL"),
                update_time,
            )
            for row in geo_heat
        ],
    )

    execute_many(
        conn,
        db_type,
        """
        INSERT INTO trend_related_queries
        (batch_id, geo_name, geo_code, timeframe_name, timeframe_param, category, keyword,
         list_type, rank_no, query, trend_value, formatted_value, link, source_url, update_time)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                batch_id,
                row.get("区域"),
                row.get("区域代码"),
                row.get("时间维度"),
                row.get("时间参数"),
                row.get("分类"),
                row.get("关键词"),
                row.get("列表类型"),
                to_int(row.get("排名")),
                row.get("查询词"),
                to_int(row.get("热度值")),
                row.get("格式化热度"),
                row.get("链接"),
                row.get("来源URL"),
                update_time,
            )
            for row in related_queries
        ],
    )

    execute_many(
        conn,
        db_type,
        """
        INSERT INTO trend_related_topics
        (batch_id, geo_name, geo_code, timeframe_name, timeframe_param, category, keyword,
         list_type, rank_no, topic, topic_type, trend_value, formatted_value, link, source_url, update_time)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                batch_id,
                row.get("区域"),
                row.get("区域代码"),
                row.get("时间维度"),
                row.get("时间参数"),
                row.get("分类"),
                row.get("关键词"),
                row.get("列表类型"),
                to_int(row.get("排名")),
                row.get("主题"),
                row.get("主题类型"),
                to_int(row.get("热度值")),
                row.get("格式化热度"),
                row.get("链接"),
                row.get("来源URL"),
                update_time,
            )
            for row in related_topics
        ],
    )

    execute(
        conn,
        db_type,
        "UPDATE crawl_runs SET status = ?, finished_at = ?, update_time = ? WHERE batch_id = ?",
        ["success", update_time, update_time, batch_id],
    )
    conn.commit()
    return {
        "batch_id": batch_id,
        "workbook": str(workbook_path),
        "timeseries": len(timeseries),
        "geo_heat": len(geo_heat),
        "related_queries": len(related_queries),
        "related_topics": len(related_topics),
        "update_time": update_time,
    }


def main():
    config = load_env(ENV_PATH)
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else latest_workbook()
    conn, db_type = connect(config)
    try:
        result = import_workbook(conn, db_type, workbook_path)
    finally:
        conn.close()
    print(result)


if __name__ == "__main__":
    main()
