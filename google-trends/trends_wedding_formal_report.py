import json
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CATEGORY = 988
PROPERTY = ""
HL = "zh-CN"
TZ = "-480"

KEYWORDS = [
    ("Wedding", "Wedding"),
    ("Wedding", "Bride"),
    ("Wedding", "Bridesmaid"),
    ("Wedding", "mother of the bride"),
    ("Wedding", "mother of the groom"),
    ("Wedding", "wedding guest"),
    ("Formal", "Military ball"),
    ("Formal", "Black tie"),
    ("Formal", "Gala"),
    ("Formal", "Evening"),
    ("Formal", "Cruise"),
    ("Semi formal", "Party"),
    ("Semi formal", "Cocktail"),
    ("Semi formal", "Outdoor wedding"),
    ("Semi formal", "Garden wedding"),
    ("Semi formal", "Birthday"),
]

GEOS = [("US", "美国"), ("GB", "英国")]
TIMES = [("today 12-m", "过去12个月"), ("today 1-m", "过去30天")]
CACHE_PATH = Path.cwd() / "google_trends_wedding_formal_single_keyword_cache.jsonl"


def clean_google_json(text):
    return json.loads(text[5:] if text.startswith(")]}'") else text)


def make_session():
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/126.0.0.0 Safari/537.36"
            ),
            "Accept": "application/json, text/plain, */*",
            "Accept-Language": "zh-CN,zh;q=0.9,en-US;q=0.8,en;q=0.7",
            "Referer": "https://trends.google.com/trends/explore",
        }
    )
    session.get("https://trends.google.com/", timeout=30)
    session.get("https://trends.google.com/trends/", timeout=30)
    return session


def explore(session, keywords, geo, timeframe):
    req = {
        "comparisonItem": [
            {"keyword": keyword, "geo": geo, "time": timeframe} for keyword in keywords
        ],
        "category": CATEGORY,
        "property": PROPERTY,
    }
    response = session.get(
        "https://trends.google.com/trends/api/explore",
        params={"hl": HL, "tz": TZ, "req": json.dumps(req, separators=(",", ":"))},
        timeout=30,
    )
    for delay in (20, 60, 120):
        if response.status_code != 429:
            break
        time.sleep(delay)
        response = session.get(
            "https://trends.google.com/trends/api/explore",
            params={"hl": HL, "tz": TZ, "req": json.dumps(req, separators=(",", ":"))},
            timeout=30,
        )
    response.raise_for_status()
    return clean_google_json(response.text)


def widget_by_id(explore_json, widget_id):
    for widget in explore_json.get("widgets", []):
        if widget.get("id") == widget_id or str(widget.get("id", "")).startswith(widget_id):
            return widget
        if widget_id == "RELATED_TOPICS" and "主题" in str(widget.get("title", "")):
            return widget
        if widget_id == "RELATED_QUERIES" and "查询" in str(widget.get("title", "")):
            return widget
    return None


def fetch_widget(session, widget, endpoint):
    token = widget["token"]
    request = widget["request"]
    response = session.get(
        f"https://trends.google.com/trends/api/widgetdata/{endpoint}",
        params={
            "hl": HL,
            "tz": TZ,
            "req": json.dumps(request, separators=(",", ":")),
            "token": token,
        },
        timeout=30,
    )
    for delay in (20, 60, 120):
        if response.status_code != 429:
            break
        time.sleep(delay)
        response = session.get(
            f"https://trends.google.com/trends/api/widgetdata/{endpoint}",
            params={
                "hl": HL,
                "tz": TZ,
                "req": json.dumps(request, separators=(",", ":")),
                "token": token,
            },
            timeout=30,
        )
    response.raise_for_status()
    return clean_google_json(response.text)


def trend_url(keyword, geo, timeframe):
    return (
        "https://trends.google.com/trends/explore"
        f"?cat={CATEGORY}&date={quote(timeframe)}&geo={geo}&gprop=&q={quote(keyword)}&hl=zh-CN"
    )


def batched(items, size):
    for idx in range(0, len(items), size):
        yield items[idx : idx + size]


def normalize_entry(item):
    topic = item.get("topic") or {}
    query = item.get("query") or item.get("title") or {}
    if not isinstance(topic, dict):
        topic = {"title": str(topic)}
    if not isinstance(query, dict):
        query = {"title": str(query)}
    title = topic.get("title") or query.get("title") or item.get("title") or item.get("query")
    if isinstance(title, dict):
        title = title.get("title")
    return {
        "title": title or "",
        "type": topic.get("type") or query.get("type") or item.get("type") or "",
        "value": item.get("value", ""),
        "formattedValue": item.get("formattedValue", ""),
        "link": item.get("link", ""),
    }


def fetch_single_keyword_related(session, category_name, keyword, geo, geo_name, timeframe, timeframe_name):
    if keyword == "Military ball" and geo != "US":
        return [], []

    data = explore(session, [keyword], geo, timeframe)
    related_topics = widget_by_id(data, "RELATED_TOPICS")
    related_queries = widget_by_id(data, "RELATED_QUERIES")

    topic_rows = []
    query_rows = []

    if related_topics:
        topics = fetch_widget(session, related_topics, "relatedsearches")
        ranked = topics.get("default", {}).get("rankedList", [])
        for list_name, items in (("热门", ranked[0].get("rankedKeyword", []) if len(ranked) > 0 else []), ("飙升", ranked[1].get("rankedKeyword", []) if len(ranked) > 1 else [])):
            for rank, item in enumerate(items, 1):
                entry = normalize_entry(item)
                topic_rows.append(
                    [
                        geo_name,
                        geo,
                        timeframe_name,
                        timeframe,
                        category_name,
                        keyword,
                        list_name,
                        rank,
                        entry["title"],
                        entry["type"],
                        entry["value"],
                        entry["formattedValue"],
                        entry["link"],
                        trend_url(keyword, geo, timeframe),
                    ]
                )

    if related_queries:
        queries = fetch_widget(session, related_queries, "relatedsearches")
        ranked = queries.get("default", {}).get("rankedList", [])
        for list_name, items in (("热门", ranked[0].get("rankedKeyword", []) if len(ranked) > 0 else []), ("飙升", ranked[1].get("rankedKeyword", []) if len(ranked) > 1 else [])):
            for rank, item in enumerate(items, 1):
                entry = normalize_entry(item)
                query_rows.append(
                    [
                        geo_name,
                        geo,
                        timeframe_name,
                        timeframe,
                        category_name,
                        keyword,
                        list_name,
                        rank,
                        entry["title"],
                        entry["value"],
                        entry["formattedValue"],
                        entry["link"],
                        trend_url(keyword, geo, timeframe),
                    ]
                )

    time.sleep(0.8)
    return topic_rows, query_rows


def collect():
    session = make_session()
    keyword_names = [keyword for _, keyword in KEYWORDS]
    categories = {keyword: category for category, keyword in KEYWORDS}

    interest_rows = []
    heat_rows = []
    topic_rows = []
    query_rows = []
    completed = set()
    if CACHE_PATH.exists():
        for line in CACHE_PATH.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            record = json.loads(line)
            if record["key"][0] == "single" and not record.get("topics"):
                continue
            if record["key"][0] == "metrics_single_keyword" and not record.get("interest"):
                continue
            completed.add(tuple(record["key"]))
            interest_rows.extend(record.get("interest", []))
            heat_rows.extend(record.get("heat", []))
            topic_rows.extend(record.get("topics", []))
            query_rows.extend(record.get("queries", []))

    for geo, geo_name in GEOS:
        for timeframe, timeframe_name in TIMES:
            comparable = [kw for kw in keyword_names if not (kw == "Military ball" and geo != "US")]
            for keyword in comparable:
                metric_key = ("metrics_single_keyword", geo, timeframe, keyword)
                if metric_key in completed:
                    continue
                metric_interest = []
                metric_heat = []
                data = explore(session, [keyword], geo, timeframe)

                interest_widget = widget_by_id(data, "TIMESERIES")
                if interest_widget:
                    timeline = fetch_widget(session, interest_widget, "multiline")
                    timeline_data = timeline.get("default", {}).get("timelineData", [])
                    if not timeline_data:
                        metric_interest.append(
                            [
                                geo_name,
                                geo,
                                timeframe_name,
                                timeframe,
                                "Google未返回时间序列",
                                "",
                                categories[keyword],
                                keyword,
                                "",
                                "",
                                trend_url(keyword, geo, timeframe),
                            ]
                        )
                    for point in timeline_data:
                        values = point.get("value", [])
                        formatted = point.get("formattedValue", [])
                        metric_interest.append(
                            [
                                geo_name,
                                geo,
                                timeframe_name,
                                timeframe,
                                point.get("formattedTime", ""),
                                point.get("time", ""),
                                categories[keyword],
                                keyword,
                                values[0] if values else "",
                                formatted[0] if formatted else "",
                                trend_url(keyword, geo, timeframe),
                            ]
                        )

                geo_widget = widget_by_id(data, "GEO_MAP")
                if geo_widget:
                    geo_data = fetch_widget(session, geo_widget, "comparedgeo")
                    for item in geo_data.get("default", {}).get("geoMapData", []):
                        values = item.get("value", [])
                        formatted = item.get("formattedValue", [])
                        metric_heat.append(
                            [
                                geo_name,
                                geo,
                                timeframe_name,
                                timeframe,
                                categories[keyword],
                                keyword,
                                item.get("geoName", ""),
                                item.get("geoCode", ""),
                                values[0] if values else "",
                                formatted[0] if formatted else "",
                                trend_url(keyword, geo, timeframe),
                            ]
                        )
                interest_rows.extend(metric_interest)
                heat_rows.extend(metric_heat)
                with CACHE_PATH.open("a", encoding="utf-8") as fh:
                    fh.write(json.dumps({"key": metric_key, "interest": metric_interest, "heat": metric_heat}, ensure_ascii=False) + "\n")
                completed.add(metric_key)
                time.sleep(2.0)

            if geo == "GB":
                na_key = ("na", geo, timeframe, "Military ball")
                if na_key not in completed:
                    na_row = [geo_name, geo, timeframe_name, timeframe, "Formal", "Military ball", "不适用", "", "仅限美国", "", trend_url("Military ball", "US", timeframe)]
                    heat_rows.append(na_row)
                    with CACHE_PATH.open("a", encoding="utf-8") as fh:
                        fh.write(json.dumps({"key": na_key, "heat": [na_row]}, ensure_ascii=False) + "\n")
                    completed.add(na_key)

            for category_name, keyword in KEYWORDS:
                single_key = ("single", geo, timeframe, keyword)
                if single_key in completed:
                    continue
                topics, queries = fetch_single_keyword_related(
                    session, category_name, keyword, geo, geo_name, timeframe, timeframe_name
                )
                topic_rows.extend(topics)
                query_rows.extend(queries)
                with CACHE_PATH.open("a", encoding="utf-8") as fh:
                    fh.write(json.dumps({"key": single_key, "topics": topics, "queries": queries}, ensure_ascii=False) + "\n")
                completed.add(single_key)

    return interest_rows, heat_rows, topic_rows, query_rows


def write_workbook(path, tables):
    wb = Workbook()
    wb.remove(wb.active)
    specs = [
        (
            "热度随时间变化趋势",
            ["区域", "区域代码", "时间维度", "时间参数", "日期", "时间戳", "分类", "关键词", "热度值", "格式化热度", "来源URL"],
            tables[0],
        ),
        (
            "搜索热度",
            ["区域", "区域代码", "时间维度", "时间参数", "子区域", "子区域代码", "分类", "关键词", "热度值", "格式化热度", "来源URL"],
            [[r[0], r[1], r[2], r[3], r[6], r[7], r[4], r[5], r[8], r[9], r[10]] for r in tables[1]],
        ),
        (
            "相关主题",
            ["区域", "区域代码", "时间维度", "时间参数", "分类", "关键词", "列表类型", "排名", "主题", "主题类型", "热度值", "格式化热度", "链接", "来源URL"],
            tables[2],
        ),
        (
            "相关查询",
            ["区域", "区域代码", "时间维度", "时间参数", "分类", "关键词", "列表类型", "排名", "查询词", "热度值", "格式化热度", "链接", "来源URL"],
            tables[3],
        ),
    ]
    for title, headers, rows in specs:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            max_len = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
            ws.column_dimensions[get_column_letter(column[0].column)].width = max_len
    wb.save(path)


if __name__ == "__main__":
    output = Path.cwd() / f"google_trends_wedding_formal_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    tables = collect()
    write_workbook(output, tables)
    print(output)
    print("rows", [len(table) for table in tables])
